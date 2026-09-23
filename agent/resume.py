"""Extract resume text from an in-memory upload; never persist the source file."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_RESUME_CHARS = 20_000


def validate_resume_text(text: str) -> str:
    cleaned = text.replace("\x00", "").strip()
    if not cleaned:
        raise ValueError("没有提取到简历文字。扫描版 PDF 请先做 OCR，或直接粘贴文字。")
    if len(cleaned) > MAX_RESUME_CHARS:
        raise ValueError(f"简历文字超过 {MAX_RESUME_CHARS} 字，请精简后再试。")
    return cleaned


def parse_resume(filename: str, data: bytes) -> str:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("简历文件不能超过 5 MB。")
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = data.decode("gb18030")
            except UnicodeDecodeError as exc:
                raise ValueError("TXT 文件编码无法识别，请保存为 UTF-8。") from exc
    elif suffix == ".pdf":
        if not data.startswith(b"%PDF"):
            raise ValueError("文件内容不是有效的 PDF。")
        from pypdf import PdfReader

        try:
            reader = PdfReader(BytesIO(data))
            if len(reader.pages) > 20:
                raise ValueError("PDF 超过 20 页，请上传精简版简历。")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("PDF 解析失败；可尝试粘贴简历文字。") from exc
    elif suffix == ".docx":
        if not data.startswith(b"PK"):
            raise ValueError("文件内容不是有效的 DOCX。")
        from docx import Document

        try:
            with ZipFile(BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError("DOCX 解压后的内容过大，请上传精简版简历。")
            doc = Document(BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
            text = "\n".join(parts)
        except ValueError:
            raise
        except BadZipFile as exc:
            raise ValueError("文件内容不是有效的 DOCX。") from exc
        except Exception as exc:
            raise ValueError("DOCX 解析失败；可尝试粘贴简历文字。") from exc
    else:
        raise ValueError("只支持 PDF、DOCX、TXT 简历。")
    return validate_resume_text(text)
